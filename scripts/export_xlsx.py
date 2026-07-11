#!/usr/bin/env python3
"""마스터 CSV → 연차 성과표 형식 xlsx 출력 (제출용 근사 초안).

사용법:
  /usr/bin/python3 scripts/export_xlsx.py 2025                  # 2025년 성과만
  /usr/bin/python3 scripts/export_xlsx.py all -o /tmp/전체.xlsx  # 전체 연도

시트: '논문,학술대회' / '특허' / 'SW' / '홍보' — 공식 성과표 컬럼 구조를 본뜸.
마스터에 없는 항목(초록·ISSN·기여율·IF 등)은 '-' 로 채우므로 제출 전 보완 필요.

의존성: openpyxl (requirements: pip install openpyxl — 이 스크립트만 필요).
"""
import argparse
import sys

import lba_master as M

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
except ImportError:
    sys.exit("openpyxl이 필요합니다: /usr/bin/python3 -m pip install openpyxl")

# 과제 공통 상수
EZONE = "2022-0-00951"
IRIS = "RS-2022-II220951"
PROJECT = "스스로 불확실성을 자각하며 질문하면서 성장하는 에이전트 기술 개발"
LEAD_ORG = "한국전자통신연구원"
PI = "장민수"

KIND_TO_GUBUN = {"SCI(E) 저널": "SCI", "비SCI(E) 저널": "비SCI",
                 "국제학술대회": "학술대회", "국내학술대회": "학술대회"}
COUNTRY_CODE = {"대한민국": "KR", "미국": "US", "일본": "JP", "PCT": "PCT"}

HEADERS = {
    "논문,학술대회": ["순번", "성과발생연도", "SCI,\n학술대회\n구분", "EZONE\n과제번호",
                 "IRIS\n과제번호", "과제명", "주관기관", "연구책임자\n(주관기관)",
                 "논문\n등록기관", "학술지명\n학술대회명*", "상위\n10위 학회", "논문명*",
                 "ISBN 또는 ISSN*", "주저자명(제1저자)*", "공동저자명", "볼륨번호*",
                 "학술지출판일자", "학술대회\n발표일자", "초록", "기여율", "IF"],
    "특허": ["순번", "성과발생연도", "EZONE\n과제번호", "IRIS\n과제번호", "과제명",
           "주관기관", "연구책임자\n(주관기관)", "출원(등록)기관*", "출원(등록)국*",
           "국내외 구분", "★출원(등록)구분*", "출원(등록)번호*",
           "발명(고안, 디자인)의 명칭*", "출원(등록)일*", "기여율"],
    "SW": ["순번", "성과발생연도", "EZONE\n과제번호", "IRIS\n과제번호", "과제명",
          "주관기관", "연구책임자\n(주관기관)", "성과발생 기관", "저작물 명칭",
          "저작물 종류", "저작자", "등록월일"],
    "홍보": ["순번", "성과발생연도", "EZONE\n과제번호", "IRIS\n과제번호", "과제명",
           "주관기관", "연구책임자\n(주관기관)", "성과발생 기관", "유형", "언론사",
           "방식", "제목", "일자"],
}

COMMON = [EZONE, IRIS, PROJECT, LEAD_ORG, PI]


def select_year(rows, year):
    if year == "all":
        return sorted(rows, key=lambda r: (-int(r["year"]), r["_line"]))
    return [r for r in rows if r["year"] == year]


def paper_rows(rows):
    kind_order = {k: i for i, k in enumerate(M.KINDS)}
    rows = sorted(rows, key=lambda r: (kind_order[M.split_multi(r["kind"])[0]],
                                       -int(r["year"]), r["_line"]))
    out = []
    for i, r in enumerate(rows, 1):
        kinds = M.split_multi(r["kind"])
        authors = M.split_multi(r["authors"])
        out.append([i, int(r["year"]), KIND_TO_GUBUN[kinds[0]], *COMMON,
                    M.split_multi(r["orgs"])[0], "; ".join(M.split_multi(r["venue"])),
                    "X", r["title"], "-", authors[0] if authors else "-",
                    ", ".join(authors[1:]), "-", "-", "-", "-", "-", "-"])
    return out


def patent_rows(rows):
    out = []
    i = 0
    for r in rows:
        countries = M.split_multi(r["countries"])
        numbers = M.split_multi(r["numbers"])
        dates = M.split_multi(r["dates"])
        org = M.split_multi(r["orgs"])[0]
        for country, number, date in zip(countries, numbers, dates):
            i += 1
            gubun = "등록" if "등록" in country else "출원"
            nation = country.replace("출원", "").replace("등록", "").strip()
            code = COUNTRY_CODE.get(nation, nation)
            inout = "01.국내" if code == "KR" else "02.국외"
            out.append([i, int(r["year"]), *COMMON, org, code, inout, gubun,
                        number, r["title"], date, "-"])
    return out


def sw_rows(rows):
    return [[i, int(r["year"]), *COMMON, r["org"], r["title"], "프로그램",
             r["author"], r.get("status") or r["date"]]
            for i, r in enumerate(rows, 1)]


def publicity_rows(rows):
    return [[i, int(r["year"]), *COMMON, r["org"], r["type"], r["outlet"],
             r["mode"], r["title"], r["date"]]
            for i, r in enumerate(rows, 1)]


def style_sheet(ws, widths):
    thin = Border(*[Side(style="thin")] * 4)
    fill = PatternFill("solid", fgColor="DDEBF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   wrap_text=True)
    for row in ws.iter_rows():
        for cell in row:
            cell.border = thin
    for col, w in zip("ABCDEFGHIJKLMNOPQRSTU", widths):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"


def main():
    ap = argparse.ArgumentParser(description="LBA 성과표 xlsx 출력")
    ap.add_argument("year", help="대상 연도(예: 2025) 또는 all")
    ap.add_argument("-o", "--out", default=None, help="출력 파일 경로")
    args = ap.parse_args()
    if args.year != "all" and not (args.year.isdigit() and
                                   M.YEAR_MIN <= int(args.year) <= M.YEAR_MAX):
        ap.error(f"연도는 {M.YEAR_MIN}~{M.YEAR_MAX} 또는 all")

    master = M.load_master()
    out_path = args.out or f"LBA_성과표_{args.year}.xlsx"

    wb = Workbook()
    wb.remove(wb.active)

    sheets = [
        ("논문,학술대회", paper_rows(select_year(master["papers"][1], args.year)),
         [5, 9, 9, 13, 16, 30, 14, 10, 14, 30, 8, 45, 12, 14, 30, 8, 12, 12, 8, 8, 6]),
        ("특허", patent_rows(select_year(master["patents"][1], args.year)),
         [5, 9, 13, 16, 30, 14, 10, 16, 10, 10, 12, 18, 45, 12, 8]),
        ("SW", sw_rows(select_year(master["software"][1], args.year)),
         [5, 9, 13, 16, 30, 14, 10, 16, 45, 10, 20, 12]),
        ("홍보", publicity_rows(select_year(master["publicity"][1], args.year)),
         [5, 9, 13, 16, 30, 14, 10, 14, 12, 16, 12, 50, 12]),
    ]
    for name, rows, widths in sheets:
        ws = wb.create_sheet(name)
        ws.append(HEADERS[name])
        for row in rows:
            ws.append(row)
        style_sheet(ws, widths)
        print(f"{name}: {len(rows)}행")

    wb.save(out_path)
    print(f"저장: {out_path}")


if __name__ == "__main__":
    main()
